"""Bulk product import from a CSV / Excel sheet.

One row = one product. Existing products are matched by ``sku`` (preferred),
then ``slug``, then exact ``title`` — a match is updated in place, otherwise a
new product is created. Categories / brands / product types / tags named in a
row are created on the fly (scoped to the project). Images are referenced by
their file name (or numeric id) in the store's Media library
(``/admin/media/``) — upload the photos there first, then list them in the
``images`` column separated by ``|``.

CSV is always supported (stdlib). ``.xlsx`` works when ``openpyxl`` is
installed; a Google Sheet is exported with File → Download → CSV.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.text import slugify

from apps.categories.models import Category
from apps.inventory.models import InventoryItem, Warehouse

from .models import Brand, Product, ProductImage, ProductType, Tag

# Column order for the downloadable sample. ``title`` is the only required one.
TEMPLATE_HEADERS = [
    "title", "sku", "slug", "status",
    "price", "sale_price", "cost_price",
    "category", "brand", "type",
    "short_description", "description",
    "weight", "length", "width", "height",
    "barcode", "hsn_sac", "tax_class",
    "is_featured", "is_new_arrival", "is_bestseller", "search_indexed",
    "tags", "stock", "images",
]

SAMPLE_ROWS = [
    {
        "title": "Classic Cotton Tee", "sku": "TEE-001", "slug": "",
        "status": "active", "price": "799", "sale_price": "599", "cost_price": "250",
        "category": "T-Shirts", "brand": "Acme", "type": "Apparel",
        "short_description": "Soft combed-cotton crew neck.",
        "description": "<p>240 gsm combed cotton. Pre-shrunk. Unisex fit.</p>",
        "weight": "0.2", "length": "", "width": "", "height": "",
        "barcode": "8901234567890", "hsn_sac": "6109", "tax_class": "5%",
        "is_featured": "yes", "is_new_arrival": "yes", "is_bestseller": "no",
        "search_indexed": "yes",
        "tags": "cotton, everyday, unisex", "stock": "50",
        "images": "tee-front.jpg | tee-back.jpg",
    },
    {
        "title": "Enamel Travel Mug", "sku": "MUG-014", "slug": "",
        "status": "active", "price": "1249", "sale_price": "", "cost_price": "480",
        "category": "Drinkware", "brand": "Acme", "type": "Homeware",
        "short_description": "350 ml, chip-resistant enamel.",
        "description": "<p>Powder-coated steel with a rolled rim.</p>",
        "weight": "0.35", "length": "", "width": "", "height": "",
        "barcode": "", "hsn_sac": "7323", "tax_class": "12%",
        "is_featured": "no", "is_new_arrival": "no", "is_bestseller": "yes",
        "search_indexed": "yes",
        "tags": "camping, gift", "stock": "20",
        "images": "mug.jpg",
    },
]

_TRUE = {"1", "true", "yes", "y", "on", "t"}
_FALSE = {"0", "false", "no", "n", "off", "f", ""}
_STATUSES = {"active", "draft", "archived"}


class ProductImportError(Exception):
    """Bad file — nothing was imported."""


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    images_attached: int = 0
    errors: list = field(default_factory=list)   # [(line_no, message)]
    missing_images: set = field(default_factory=set)

    @property
    def ok_rows(self):
        return self.created + self.updated

    @property
    def has_errors(self):
        return bool(self.errors)


# --- sample -------------------------------------------------------

def sample_csv() -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=TEMPLATE_HEADERS)
    w.writeheader()
    for row in SAMPLE_ROWS:
        w.writerow(row)
    return buf.getvalue()


# --- reading -----------------------------------------------------

def read_table(django_file, filename: str) -> list[dict]:
    """Return a list of ``{header: value}`` dicts from an uploaded CSV/XLSX."""
    name = (filename or getattr(django_file, "name", "") or "").lower()

    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(django_file)
    return _read_csv(django_file)


def _read_csv(django_file) -> list[dict]:
    raw = django_file.read()
    if isinstance(raw, bytes):
        # Excel often saves CSV as UTF-8 with a BOM.
        raw = raw.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(raw))
    if not reader.fieldnames:
        raise ProductImportError("The file is empty.")
    return [_clean_keys(r) for r in reader]


def _read_xlsx(django_file) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ProductImportError(
            "Excel files need the openpyxl package — save the sheet as CSV instead."
        ) from exc

    wb = load_workbook(io.BytesIO(django_file.read()), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
    except StopIteration:
        raise ProductImportError("The sheet is empty.")
    out = []
    for values in rows:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        row = {}
        for i, key in enumerate(header):
            if not key:
                continue
            v = values[i] if i < len(values) else None
            row[key] = "" if v is None else str(v).strip()
        out.append(row)
    return out


def _clean_keys(row: dict) -> dict:
    return {
        (k or "").strip().lower(): (v or "").strip()
        for k, v in row.items()
        if (k or "").strip()
    }


# --- value coercion --------------------------------------------

def _decimal(value, fieldname):
    value = (value or "").strip().replace(",", "")
    if not value:
        return None
    try:
        return Decimal(value)
    except (InvalidOperation, ValueError):
        raise ValueError(f"{fieldname}: “{value}” is not a number")


def _bool(value, default=False):
    value = (value or "").strip().lower()
    if value == "":
        return default
    if value in _TRUE:
        return True
    if value in _FALSE:
        return False
    return default


def _int(value, fieldname):
    value = (value or "").strip()
    if not value:
        return None
    try:
        return int(float(value))
    except ValueError:
        raise ValueError(f"{fieldname}: “{value}” is not a whole number")


# --- import ----------------------------------------------------

def run_import(project, rows: list[dict], *, actor=None, default_status="draft") -> ImportResult:
    result = ImportResult()
    if not rows:
        raise ProductImportError("No data rows found under the header.")

    media_index = _media_index(project)
    caches = {"category": {}, "brand": {}, "type": {}}

    for n, raw in enumerate(rows, start=2):  # row 1 is the header
        row = {k: (v or "").strip() for k, v in raw.items()}
        title = row.get("title", "")
        if not title:
            if any(row.values()):
                result.errors.append((n, "title is required"))
            continue
        try:
            with transaction.atomic():
                created = _apply_row(
                    project, row, media_index, caches, result, actor=actor,
                    default_status=default_status,
                )
            if created:
                result.created += 1
            else:
                result.updated += 1
        except ValueError as exc:
            result.errors.append((n, str(exc)))
        except Exception as exc:  # noqa: BLE001 — never abort the whole file
            result.errors.append((n, f"unexpected error: {exc}"))

    return result


def _apply_row(project, row, media_index, caches, result, *, actor, default_status):
    sku = row.get("sku", "")
    slug = row.get("slug", "")
    title = row["title"]

    product = None
    if sku:
        product = Product.objects.filter(project=project, sku__iexact=sku).first()
    if product is None and slug:
        product = Product.objects.filter(project=project, slug=slugify(slug)).first()
    if product is None:
        product = Product.objects.filter(project=project, title__iexact=title).first()
    is_new = product is None
    if is_new:
        product = Product(project=project)

    product.title = title
    if slug:
        product.slug = slugify(slug)
    if sku:
        product.sku = sku

    status = (row.get("status", "") or "").lower() or default_status
    if status not in _STATUSES:
        raise ValueError(f"status: “{status}” (use active / draft / archived)")
    product.status = status

    price = _decimal(row.get("price"), "price")
    if price is not None:
        product.price = price
    elif is_new:
        product.price = Decimal("0")
    product.sale_price = _decimal(row.get("sale_price"), "sale_price")
    product.cost_price = _decimal(row.get("cost_price"), "cost_price")

    for f in ("short_description", "description", "barcode", "hsn_sac", "tax_class"):
        if row.get(f):
            setattr(product, f, row[f])

    for f in ("weight", "length", "width", "height"):
        v = _decimal(row.get(f), f)
        if v is not None:
            setattr(product, f, v)

    for f in ("is_featured", "is_new_arrival", "is_bestseller", "search_indexed"):
        if row.get(f, "") != "":
            setattr(product, f, _bool(row.get(f, ""), default=getattr(product, f)))

    if row.get("category"):
        product.category = _get_or_make(Category, project, row["category"], caches["category"])
    if row.get("brand"):
        product.brand = _get_or_make(Brand, project, row["brand"], caches["brand"])
    if row.get("type"):
        product.type = _get_or_make(ProductType, project, row["type"], caches["type"])

    product.save()

    if row.get("tags"):
        names = [t.strip() for t in row["tags"].split(",") if t.strip()]
        tags = [
            Tag.objects.get_or_create(project=project, name=name)[0]
            for name in names
        ]
        product.tags.set(tags)

    if row.get("stock") != "":
        qty = _int(row.get("stock"), "stock")
        if qty is not None:
            _set_stock(project, product, qty)

    if row.get("images"):
        result.images_attached += _attach_images(
            product, row["images"], media_index, result.missing_images
        )

    return is_new


def _get_or_make(model, project, name, cache):
    name = name.strip()
    key = name.lower()
    if key in cache:
        return cache[key]
    obj = (
        model.objects.filter(project=project, name__iexact=name).first()
        or model.objects.create(project=project, name=name)
    )
    cache[key] = obj
    return obj


def _default_warehouse(project):
    wh = (
        Warehouse.objects.filter(project=project, is_default=True).first()
        or Warehouse.objects.filter(project=project).order_by("id").first()
    )
    if wh is None:
        wh = Warehouse.objects.create(project=project, name="Main", is_default=True)
    return wh


def _set_stock(project, product, qty):
    item, _ = InventoryItem.objects.get_or_create(
        warehouse=_default_warehouse(project), product=product, variant=None,
    )
    if item.quantity != qty:
        item.quantity = qty
        item.save(update_fields=["quantity", "updated_at"])


# --- images ----------------------------------------------------

def _media_index(project):
    """Lower-cased file name (and bare stem) -> MediaAsset, plus id -> asset."""
    from apps.media.models import AssetKind, MediaAsset

    index = {}
    for a in MediaAsset.objects.filter(
        project=project, kind=AssetKind.IMAGE, trashed_at__isnull=True
    ):
        index[str(a.pk)] = a
        name = (a.original_name or a.file.name.rsplit("/", 1)[-1]).lower()
        index.setdefault(name, a)
        index.setdefault(name.rsplit(".", 1)[0], a)
    return index


def _image_hash(data: bytes) -> str:
    import hashlib

    return hashlib.md5(data).hexdigest()


def _attach_images(product, cell, media_index, missing: set) -> int:
    refs = [r.strip() for r in cell.replace("\n", "|").split("|") if r.strip()]
    if not refs:
        return 0

    from django.core.files.base import ContentFile

    existing_images = list(product.images.all())
    existing = len(existing_images)
    has_primary = any(img.is_primary for img in existing_images)
    attached = 0

    # Re-running the same (or an overlapping) import file must not re-attach an
    # image the product already has — dedupe by content, since the storage
    # backend renames files on collision so filenames alone aren't reliable.
    seen_hashes = set()
    for img in existing_images:
        try:
            with img.image.open("rb") as f:
                seen_hashes.add(_image_hash(f.read()))
        except (OSError, ValueError):
            continue

    for ref in refs:
        asset = media_index.get(ref.lower())
        if asset is None:
            missing.add(ref)
            continue
        basename = (asset.original_name or asset.file.name.rsplit("/", 1)[-1])
        try:
            data = asset.file.read()
        finally:
            asset.file.close()
        digest = _image_hash(data)
        if digest in seen_hashes:
            continue
        seen_hashes.add(digest)
        img = ProductImage(
            product=product, alt=asset.alt or product.title,
            order=existing + attached + 1,
            is_primary=(not has_primary and attached == 0),
        )
        img.image.save(basename, ContentFile(data), save=True)
        attached += 1

    return attached
